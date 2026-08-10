from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Iterable, List, Optional


def parse_roster(workbook_path: Path) -> List[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, data_only=True)
    rows: List[dict] = []
    for sheet in workbook.worksheets:
        pen = {"7315": 4, "7314": 5, "7317": 6}.get(sheet.title)
        for values in sheet.iter_rows(values_only=True):
            if len(values) < 4:
                continue
            sow, ear_tag, piglet_number, gender = values[:4]
            if not isinstance(piglet_number, (int, float)):
                continue
            piglet_id = int(piglet_number)
            if piglet_id < 1 or piglet_id > 99:
                continue
            notes = str(values[4] or "") if len(values) > 4 else ""
            rows.append(
                {
                    "pen": pen,
                    "sow_number": int(sow) if isinstance(sow, (int, float)) else sheet.title,
                    "ear_tag": int(ear_tag) if isinstance(ear_tag, (int, float)) else ear_tag,
                    "pig_id": str(piglet_id),
                    "gender": str(gender or ""),
                    "present": "not present" not in notes.lower(),
                    "notes": notes,
                }
            )
    return rows


def _write_csv(path: Path, rows: Iterable[dict], fieldnames: List[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def prepare_labeling_package(
    workbook: Path,
    marking_photos: Path,
    output_directory: Path,
    crops_directory: Optional[Path] = None,
    crops_root: Optional[Path] = None,
) -> dict:
    output_directory.mkdir(parents=True, exist_ok=True)
    roster = parse_roster(workbook)
    _write_csv(
        output_directory / "roster.csv",
        roster,
        ["pen", "sow_number", "ear_tag", "pig_id", "gender", "present", "notes"],
    )

    labels_path = output_directory / "identity_labels.csv"
    existing = {}
    if labels_path.exists():
        with labels_path.open(newline="") as handle:
            existing = {row["image_path"]: row for row in csv.DictReader(handle)}

    label_rows: List[dict] = []
    paths = sorted(marking_photos.glob("*.jpeg"))
    if crops_directory is not None:
        paths += sorted(crops_directory.rglob("*.jpg"))
    if crops_root is not None:
        for video_dir in sorted(path for path in crops_root.iterdir() if path.is_dir()):
            paths += sorted(video_dir.glob("track_*.jpg"))
    new_unlabeled = 0
    for path in paths:
        # Only marking photos like 11.jpeg / 11_1.jpeg are auto-labeled from the filename.
        # track_123.jpg must stay unlabeled for manual labeling.
        match = re.fullmatch(r"(\d+)(?:_\d+)?", path.stem)
        pig_id = match.group(1) if match else ""
        source_type = "reference" if path.parent.resolve() == marking_photos.resolve() else "track_crop"
        image_path = str(path.resolve())
        row = {
            "image_path": image_path,
            "pig_id": pig_id,
            "pen": "",
            "source_video": "" if source_type == "reference" else path.parent.name,
            "visibility": "clear" if pig_id else "",
            "status": "labeled" if pig_id else "unlabeled",
            "source_type": source_type,
        }
        if image_path in existing:
            row.update(existing[image_path])
            # Fix accidental skip tokens typed as pig_id (s/ss/skip). Do NOT treat an
            # empty pig_id as skip — that was wiping new unlabeled track crops on re-prepare.
            pig_token = str(row.get("pig_id", "")).strip()
            if pig_token and _should_skip_label_input(pig_token):
                row["pig_id"] = ""
                row["status"] = "skip"
        elif source_type == "track_crop":
            new_unlabeled += 1
        label_rows.append(row)
    _write_csv(
        labels_path,
        label_rows,
        [
            "image_path",
            "pig_id",
            "pen",
            "source_video",
            "visibility",
            "status",
            "source_type",
        ],
    )
    instructions = (
        "# Identity labeling\n\n"
        "Open `identity_labels.csv` in a spreadsheet. For every image, set `pig_id` to the "
        "painted back number, or `unknown`/`ambiguous`; set `pen`, `visibility` "
        "(`clear`, `partial`, `blurred`, `hidden`) and `status` (`labeled` or `skip`). "
        "Do not infer an ID from appearance when the painted mark is unreadable. Save as CSV.\n"
    )
    (output_directory / "LABELING.md").write_text(instructions)
    unlabeled = sum(1 for row in label_rows if row.get("status") == "unlabeled")
    by_video: dict[str, int] = {}
    for row in label_rows:
        if row.get("status") == "unlabeled" and row.get("source_video"):
            by_video[row["source_video"]] = by_video.get(row["source_video"], 0) + 1
    return {
        "roster_rows": len(roster),
        "images": len(label_rows),
        "prelabeled": sum(bool(row["pig_id"]) for row in label_rows),
        "preserved_labels": sum(
            row["image_path"] in existing and existing[row["image_path"]].get("status") == "labeled"
            for row in label_rows
        ),
        "new_unlabeled": new_unlabeled,
        "unlabeled_total": unlabeled,
        "unlabeled_by_video": by_video,
        "labels_path": str(labels_path),
    }


def interactive_label_manifest(manifest_csv: Path) -> dict:
    """Label analyze crop rows; writes manual_pig_id for comparison with model columns."""
    import cv2

    with manifest_csv.open(newline="") as handle:
        rows = list(csv.DictReader(handle))
    if not rows:
        return {"manifest": str(manifest_csv), "labeled": 0}
    fields = list(rows[0].keys())
    for column in ("manual_pig_id", "manual_status"):
        if column not in fields:
            fields.append(column)
    labeled = skipped = 0
    for index, row in enumerate(rows):
        if row.get("manual_status") in {"labeled", "skip"}:
            continue
        image_path = row.get("image_path", "")
        image = cv2.imread(image_path)
        if image is None:
            row["manual_status"] = "skip"
            row["manual_pig_id"] = ""
            skipped += 1
            _write_csv(manifest_csv, rows, fields)
            continue
        hints = []
        for key in (
            "pig_id",
            "candidate_pig_id",
            "classifier_pig_id",
            "mark_pig_id",
            "trocr_pig_id",
        ):
            value = (row.get(key) or "").strip()
            if value and value not in hints:
                hints.append(f"{key}={value}")
        hint_text = ", ".join(hints) if hints else "no model guess"
        cv2.namedWindow("Piglet identity", cv2.WINDOW_NORMAL)
        cv2.imshow("Piglet identity", image)
        cv2.waitKey(1)
        value = input(
            f"[{index + 1}/{len(rows)}] track {row.get('track_id', '?')} "
            f"({hint_text}) — manual ID (Enter/s=skip, q=quit): "
        ).strip()
        if value.lower() == "q":
            break
        if _should_skip_label_input(value):
            row["manual_status"] = "skip"
            row["manual_pig_id"] = ""
            skipped += 1
        else:
            row["manual_pig_id"] = value
            row["manual_status"] = "labeled"
            labeled += 1
        _write_csv(manifest_csv, rows, fields)
    cv2.destroyAllWindows()
    return {
        "manifest": str(manifest_csv.resolve()),
        "manual_labeled": labeled,
        "manual_skipped": skipped,
        "rows": len(rows),
    }


def score_manifest_labels(manifest_csv: Path) -> dict:
    """Compare manual_pig_id to model columns on labeled rows."""
    with manifest_csv.open(newline="") as handle:
        rows = list(csv.DictReader(handle))
    truth = [
        row
        for row in rows
        if row.get("manual_status") == "labeled"
        and (row.get("manual_pig_id") or "").strip().isdigit()
    ]
    if not truth:
        return {"manual_labeled_rows": 0}
    columns = [
        "pig_id",
        "candidate_pig_id",
        "classifier_pig_id",
        "mark_pig_id",
        "trocr_pig_id",
    ]
    metrics: dict = {"manual_labeled_rows": len(truth)}
    for column in columns:
        if column not in rows[0]:
            continue
        matches = sum(
            (row.get(column) or "").strip() == row["manual_pig_id"].strip() for row in truth
        )
        filled = sum(bool((row.get(column) or "").strip()) for row in truth)
        metrics[f"{column}_accuracy"] = matches / len(truth)
        metrics[f"{column}_coverage"] = filled / len(truth)
    return metrics


def export_manifest_labels_to_training(
    manifest_csv: Path, labels_csv: Path, source_video: str = ""
) -> dict:
    """Append manual manifest labels into identity_labels.csv for retraining."""
    with manifest_csv.open(newline="") as handle:
        manifest_rows = list(csv.DictReader(handle))
    labels_path = Path(labels_csv)
    existing = {}
    fields = [
        "image_path",
        "pig_id",
        "pen",
        "source_video",
        "visibility",
        "status",
        "source_type",
    ]
    if labels_path.exists():
        with labels_path.open(newline="") as handle:
            existing_rows = list(csv.DictReader(handle))
            if existing_rows:
                fields = list(existing_rows[0].keys())
            existing = {row["image_path"]: row for row in existing_rows}
    added = 0
    for row in manifest_rows:
        if row.get("manual_status") != "labeled":
            continue
        pig_id = (row.get("manual_pig_id") or "").strip()
        if not pig_id.isdigit():
            continue
        image_path = str(Path(row["image_path"]).resolve())
        video = source_video or row.get("source_video") or ""
        entry = {
            "image_path": image_path,
            "pig_id": pig_id,
            "pen": "",
            "source_video": video,
            "visibility": "unrated",
            "status": "labeled",
            "source_type": "track_crop",
        }
        if image_path in existing:
            entry.update(existing[image_path])
            entry["pig_id"] = pig_id
            entry["status"] = "labeled"
            entry["source_type"] = "track_crop"
        existing[image_path] = entry
        added += 1
    ordered = sorted(existing.values(), key=lambda item: item["image_path"])
    _write_csv(labels_path, ordered, fields)
    return {"labels_csv": str(labels_path), "manifest_rows_imported": added}


def _should_skip_label_input(value: str) -> bool:
    token = value.strip().lower()
    return token in {"", "s", "ss", "skip", "sk", "x"}


def reset_label_rows(
    labels_csv: Path,
    source_video: Optional[str] = None,
    from_status: str = "skip",
    to_status: str = "unlabeled",
) -> dict:
    """Re-queue rows for labeling (prepare never changes skip/labeled rows)."""
    with labels_csv.open(newline="") as handle:
        rows = list(csv.DictReader(handle))
    if not rows:
        return {"reset": 0, "labels_csv": str(labels_csv)}
    fields = list(rows[0].keys())
    reset = 0
    for row in rows:
        if row.get("status") != from_status:
            continue
        if source_video is not None and row.get("source_video") != source_video:
            continue
        if row.get("source_type") == "reference":
            continue
        row["status"] = to_status
        row["pig_id"] = ""
        reset += 1
    _write_csv(labels_csv, rows, fields)
    return {
        "labels_csv": str(labels_csv.resolve()),
        "reset": reset,
        "from_status": from_status,
        "to_status": to_status,
        "source_video": source_video or "*",
    }


def interactive_label(labels_csv: Path, ask_visibility: bool = False) -> None:
    import cv2

    with labels_csv.open(newline="") as handle:
        rows = list(csv.DictReader(handle))
        fields = list(rows[0]) if rows else []
    pending = [index for index, row in enumerate(rows) if row.get("status") not in {"labeled", "skip"}]
    if not pending:
        print(f"No unlabeled rows in {labels_csv} ({len(rows)} total). Run prepare after sample-crops.")
        return
    print(f"{len(pending)} images to label ({len(rows)} rows in CSV; reference photos may already be labeled).")
    prompt_num = 0
    for index in pending:
        row = rows[index]
        prompt_num += 1
        image = cv2.imread(row["image_path"])
        if image is None:
            row["status"] = "skip"
            _write_csv(labels_csv, rows, fields)
            continue
        cv2.namedWindow("Piglet identity", cv2.WINDOW_NORMAL)
        cv2.imshow("Piglet identity", image)
        cv2.waitKey(1)
        value = input(
            f"[{prompt_num}/{len(pending)} unlabeled, CSV row {index + 1}/{len(rows)}] "
            f"pig ID (number, Enter/s/ss=skip, q=quit): "
        ).strip()
        if value.lower() == "q":
            remaining = sum(
                1
                for item in rows
                if item.get("status") not in {"labeled", "skip"}
            )
            print(f"Stopped early; {remaining} rows still unlabeled. Re-run piglet-id label to continue.")
            break
        if _should_skip_label_input(value):
            row["status"] = "skip"
            row["pig_id"] = ""
        else:
            row["pig_id"] = value
            row["status"] = "labeled"
            row["visibility"] = (
                input("visibility [clear/partial/blurred/hidden]: ").strip()
                if ask_visibility
                else (row.get("visibility") or "unrated")
            )
        _write_csv(labels_csv, rows, fields)
    else:
        print("All rows are labeled or skipped.")
    cv2.destroyAllWindows()
